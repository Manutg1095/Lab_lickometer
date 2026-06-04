from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

def style_workbook(path):
    wb=load_workbook(path); fill=PatternFill('solid', fgColor='1F2937'); font=Font(color='FFFFFF', bold=True)
    for ws in wb.worksheets:
        ws.freeze_panes='A2'
        for cell in ws[1]: cell.fill=fill; cell.font=font; cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
        for col in ws.columns:
            letter=get_column_letter(col[0].column); max_len=max([len('' if c.value is None else str(c.value)) for c in col] + [10]); ws.column_dimensions[letter].width=min(max(10,max_len+2),45)
    wb.save(path)

def make_graphpad_grouped(per_animal, metric):
    if per_animal.empty or metric not in per_animal: return pd.DataFrame()
    rows=[]
    for exp, expdf in per_animal.groupby('experiment'):
        for session, sesdf in expdf.groupby('session'):
            row={'experiment':exp,'session':session}
            for _,r in sesdf.sort_values(['group_order','group','rat_id']).iterrows(): row[f"{r['group']}__{r['rat_id']}"]=r[metric]
            rows.append(row)
    return pd.DataFrame(rows)

def export_excel(output_path, preview_df, group_summary, long_df, per_animal, summary, hist_animal, hist_summary, licks_minute, licks_minute_summary, params):
    output_path=Path(output_path); output_path.parent.mkdir(parents=True, exist_ok=True)
    notes=pd.DataFrame([[k,v] for k,v in params.items()] + [
        ['Definición ráfaga','≥ min_licks_per_burst licks consecutivos con ILI <= burst_threshold_ms.'],
        ['IBI corregido','inicio de la siguiente ráfaga válida − fin de la ráfaga previa.'],
        ['Histograma','frecuencia relativa por animal; después media grupal ± SEM.']], columns=['campo','valor'])
    gp={'GraphPad_latencia_lick':'latencia_primer_lick_ms','GraphPad_latencia_burst':'latencia_primer_burst_ms','GraphPad_ILI_60_220':'media_ILI_60_220_ms','GraphPad_ILI_220_330':'media_ILI_220_330_ms','GraphPad_tamano_rafaga':'tamano_rafaga_media_licks','GraphPad_IBI':'IBI_corregido_media_ms'}
    with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
        notes.to_excel(writer, sheet_name='README', index=False); preview_df.to_excel(writer, sheet_name='Metadata_detectada', index=False); group_summary.to_excel(writer, sheet_name='Grupos_detectados', index=False)
        per_animal.to_excel(writer, sheet_name='Por_animal', index=False); summary.to_excel(writer, sheet_name='Resumen_grupo_sesion', index=False); licks_minute.to_excel(writer, sheet_name='Licks_por_minuto', index=False); licks_minute_summary.to_excel(writer, sheet_name='Licks_min_resumen', index=False); hist_animal.to_excel(writer, sheet_name='Histograma_S2_animal', index=False); hist_summary.to_excel(writer, sheet_name='Histograma_S2_resumen', index=False)
        cols=[c for c in ['sheet','prefix','experiment','rat_id','session','group','group_code','t_rel_ms','delta_ms','indice_lick'] if c in long_df]
        long_df[cols].to_excel(writer, sheet_name='Datos_largos', index=False)
        for sh,m in gp.items(): make_graphpad_grouped(per_animal,m).to_excel(writer, sheet_name=sh[:31], index=False)
    style_workbook(output_path)
